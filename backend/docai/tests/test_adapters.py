from __future__ import annotations

import json
import zipfile
from types import SimpleNamespace

import pytest

from docai.adapters import azure_identity
from docai.adapters.layout import azure_di
from docai.adapters.layout.azure_di import AzureDocumentIntelligenceLayout
from docai.adapters.layout.base import get_layout_provider_for_format
from docai.adapters.layout.excel import _col_letter, excel_layout, inspect_xlsx_safety
from docai.adapters.layout.fixture import FixtureLayout
from docai.adapters.layout.plain_text import text_layout
from docai.adapters.llm import azure_openai
from docai.adapters.llm.azure_openai import AzureOpenAILangChainLLM
from docai.adapters.llm.base import LLMCall
from docai.exceptions import (
    CorruptFile,
    IntegrationError,
    InvalidModelOutput,
    ThrottledUpstream,
    UnsafeWorkbook,
)
from docai.schemas.llm import ClassificationOut


def _span(offset: int, length: int):
    return SimpleNamespace(offset=offset, length=length)


def test_layout_provider_resolves_native_and_configured_formats():
    assert get_layout_provider_for_format("xlsx", "fixture").key == "excel"
    assert get_layout_provider_for_format("txt", "fixture").key == "plain_text"
    assert get_layout_provider_for_format("pdf", "fixture").key == "fixture"


def _azure_result():
    content = "Header Total 100"
    region = SimpleNamespace(page_number=1)
    return SimpleNamespace(
        content=content,
        model_id="prebuilt-layout",
        paragraphs=[
            SimpleNamespace(
                bounding_regions=[region], content="Header", role="title", spans=[_span(0, 6)]
            ),
            SimpleNamespace(bounding_regions=[], content="orphan", role=None, spans=[]),
        ],
        tables=[
            SimpleNamespace(
                bounding_regions=[region],
                row_count=1,
                column_count=1,
                cells=[
                    SimpleNamespace(
                        row_index=0,
                        column_index=0,
                        row_span=None,
                        column_span=None,
                        kind=None,
                        content="100",
                        spans=[_span(13, 3)],
                    )
                ],
            )
        ],
        pages=[
            SimpleNamespace(
                page_number=1,
                width=10,
                height=20,
                unit="inch",
                angle=0,
                spans=[_span(0, len(content))],
                words=[
                    SimpleNamespace(
                        content="Header",
                        polygon=[1, 2, 4, 2, 4, 4, 1, 4],
                        span=_span(0, 6),
                        confidence=0.99,
                    ),
                    SimpleNamespace(content="100", polygon=[], span=None, confidence=None),
                ],
                lines=[
                    SimpleNamespace(
                        content="Header Total",
                        polygon=[1, 2, 8, 2, 8, 4, 1, 4],
                        spans=[_span(0, 12)],
                    )
                ],
                selection_marks=[
                    SimpleNamespace(
                        state="selected", polygon=[8, 2, 9, 2, 9, 3, 8, 3], span=_span(12, 1)
                    )
                ],
            )
        ],
        sections=[SimpleNamespace(spans=[_span(0, 16)], elements=["/paragraphs/0"])],
    )


def test_azure_error_mapping_and_retry_policy(monkeypatch, settings):
    rate_limit = type("ServiceError", (Exception,), {"status_code": 429})()
    auth = type("AuthenticationError", (Exception,), {})()
    gateway_timeout = type(
        "ServiceError", (Exception,), {"response": SimpleNamespace(status_code=504)}
    )()

    assert isinstance(azure_identity.sanitize_azure_error(rate_limit), ThrottledUpstream)
    auth_error = azure_identity.sanitize_azure_error(auth)
    assert auth_error.error_code == "AZURE_AUTH_FAILED"
    assert auth_error.retryable is False
    assert azure_identity.sanitize_azure_error(gateway_timeout).error_code == "AZURE_TIMEOUT"
    assert azure_identity.sanitize_azure_error(RuntimeError()).error_code == "AZURE_ERROR"

    calls = 0
    delays: list[float] = []

    def transient_call():
        nonlocal calls
        calls += 1
        if calls < 3:
            raise TimeoutError
        return "ready"

    monkeypatch.setattr(azure_identity.time, "sleep", delays.append)
    assert azure_identity.with_retries(transient_call, max_retries=2, base_delay=0.25) == "ready"
    assert delays == [0.25, 0.5]

    with pytest.raises(IntegrationError, match="authentication failed"):
        azure_identity.with_retries(lambda: (_ for _ in ()).throw(auth), max_retries=3)
    with pytest.raises(ValueError, match="zero or greater"):
        azure_identity.with_retries(lambda: "unused", max_retries=-1)

    settings.DOCAI = {**settings.DOCAI, "AZURE_TIMEOUT_S": 17}
    assert azure_identity.azure_settings()["AZURE_TIMEOUT_S"] == 17


def test_azure_credential_and_token_provider_are_cached(monkeypatch):
    import azure.identity

    created = []

    class FakeCredential:
        def __init__(self, **kwargs):
            created.append(kwargs)

        def get_token(self, scope):
            return SimpleNamespace(token=f"token-for:{scope}")

    azure_identity.credential.cache_clear()
    monkeypatch.setattr(azure.identity, "DefaultAzureCredential", FakeCredential)
    first = azure_identity.credential()

    assert azure_identity.credential() is first
    assert created == [{"exclude_interactive_browser_credential": True}]
    assert azure_identity.token_provider("scope")() == "token-for:scope"
    azure_identity.credential.cache_clear()


def test_azure_document_intelligence_normalizes_source_coordinates():
    adapter = object.__new__(AzureDocumentIntelligenceLayout)
    adapter.api_version = "2024-11-30"

    layout = adapter.normalize(_azure_result(), document_id="doc-1", source_format="pdf")

    page = layout.pages[0]
    assert layout.model_id == "prebuilt-layout"
    assert page.content == "Header Total 100"
    assert page.words[0].polygon == [0.1, 0.1, 0.4, 0.1, 0.4, 0.2, 0.1, 0.2]
    assert page.words[0].span is not None
    assert page.words[0].span.offset == 0
    assert page.words[1].span is None
    assert page.paragraphs[0].role == "title"
    assert page.tables[0].cells[0].text == "100"
    assert page.tables[0].cells[0].row_span == 1
    assert page.selection_marks[0].state == "selected"
    assert page.reading_order == ["p1:para0", "p1:para1", "p1:t0"]
    assert layout.sections == [
        {"spans": [{"offset": 0, "length": 16}], "elements": ["/paragraphs/0"]}
    ]
    assert azure_di._norm_poly([], 10, 20) == []
    assert azure_di._norm_poly([1, 2], 0, 20) == []
    assert azure_di._span([]) is None


def test_azure_document_intelligence_configuration_and_analysis(tmp_path, monkeypatch, settings):
    settings.DOCAI = {
        **settings.DOCAI,
        "AZURE_DI_ENDPOINT": "https://example.cognitiveservices.azure.com",
        "AZURE_DI_API_VERSION": "2024-11-30",
        "AZURE_TIMEOUT_S": 2,
    }
    adapter = AzureDocumentIntelligenceLayout()
    source = tmp_path / "scan.pdf"
    source.write_bytes(b"pdf")
    result = _azure_result()
    calls = []

    class Poller:
        def result(self, *, timeout):
            assert timeout == 20
            return result

    class Client:
        def begin_analyze_document(self, model, request, *, features):
            calls.append((model, request.bytes_source, features))
            return Poller()

    monkeypatch.setattr(adapter, "_client", lambda: Client())
    monkeypatch.setattr(azure_di, "with_retries", lambda fn: fn())

    layout = adapter.analyze(source, document_id="doc-2", source_format="pdf")
    assert layout.pages[0].content == "Header Total 100"
    assert calls == [("prebuilt-layout", b"pdf", ["keyValuePairs"])]

    settings.DOCAI = {**settings.DOCAI, "AZURE_DI_ENDPOINT": ""}
    with pytest.raises(RuntimeError, match="AZURE_DI_ENDPOINT"):
        AzureDocumentIntelligenceLayout()


def test_azure_openai_adapter_returns_auditable_structured_result(monkeypatch, settings):
    settings.DOCAI = {
        **settings.DOCAI,
        "AZURE_OPENAI_ENDPOINT": "https://example.openai.azure.com",
        "AZURE_OPENAI_API_VERSION": "2024-10-21",
        "AZURE_OPENAI_DEPLOYMENT": "classification-model",
        "AZURE_TIMEOUT_S": 9,
    }
    adapter = AzureOpenAILangChainLLM(parameters={"max_retries": 0, "temperature": 0.1})

    class StructuredModel:
        def invoke(self, messages):
            assert messages == [("system", "Classify"), ("user", "Form W-2")]
            return {
                "raw": SimpleNamespace(content="raw model response", additional_kwargs={}),
                "parsed": {"category": "w2", "confidence": 0.98},
                "parsing_error": None,
            }

    class Model:
        def with_structured_output(self, schema, *, include_raw):
            assert schema is ClassificationOut
            assert include_raw is True
            return StructuredModel()

    monkeypatch.setattr(adapter, "_model", lambda deployment, params: Model())
    monkeypatch.setattr(azure_openai, "with_retries", lambda fn, max_retries: fn())
    call = LLMCall(
        system="Classify",
        user="Form W-2",
        schema=ClassificationOut,
        prompt_name="classify",
        prompt_version=2,
        schema_name="classification",
        schema_version=1,
    )

    result = adapter.invoke(call)
    assert result.parsed.category == "w2"
    assert result.raw_response == "raw model response"
    assert result.model_deployment == "classification-model"
    assert result.parameters["temperature"] == 0.1
    assert "timeout_s" not in result.parameters
    assert result.prompt_version == 2
    assert result.input_chars == len("ClassifyForm W-2")


@pytest.mark.parametrize(
    "output",
    [
        {"raw": None, "parsed": None, "parsing_error": ValueError("bad JSON")},
        {
            "raw": SimpleNamespace(content="", additional_kwargs={"refusal": "no"}),
            "parsed": {"category": "", "confidence": 2},
            "parsing_error": None,
        },
    ],
)
def test_azure_openai_adapter_rejects_invalid_model_output(output, monkeypatch, settings):
    settings.DOCAI = {
        **settings.DOCAI,
        "AZURE_OPENAI_ENDPOINT": "https://example.openai.azure.com",
    }
    adapter = AzureOpenAILangChainLLM()

    class Model:
        def with_structured_output(self, schema, *, include_raw):
            return SimpleNamespace(invoke=lambda messages: output)

    monkeypatch.setattr(adapter, "_model", lambda deployment, params: Model())
    monkeypatch.setattr(azure_openai, "with_retries", lambda fn, max_retries: fn())

    with pytest.raises(InvalidModelOutput) as exc:
        adapter.invoke(LLMCall(system="s", user="u", schema=ClassificationOut))
    assert isinstance(exc.value.errors, dict)
    assert exc.value.errors["schema"] == "ClassificationOut"


def test_local_text_and_fixture_adapters(tmp_path):
    text_file = tmp_path / "notes.txt"
    text_file.write_text("Account 123\n\nBalance 42", encoding="utf-8")

    plain = text_layout(text_file, document_id="plain-doc")
    assert plain.service == "plain_text"
    assert [line.text for line in plain.pages[0].lines] == ["Account 123", "Balance 42"]
    assert plain.pages[0].lines[0].word_ids == ["p1:w0", "p1:w1"]

    generated = FixtureLayout().analyze(text_file, document_id="fixture-doc", source_format="txt")
    assert generated.pages[0].paragraphs[0].text == "Account 123"
    assert generated.pages[0].has_text_layer is True

    sidecar = text_file.with_suffix(".txt.layout.json")
    sidecar.write_text(
        json.dumps(
            {
                "document_id": "replaced",
                "source_format": "txt",
                "service": "fixture",
                "units": [],
            }
        ),
        encoding="utf-8",
    )
    loaded = FixtureLayout().analyze(text_file, document_id="sidecar-doc", source_format="txt")
    assert loaded.document_id == "sidecar-doc"
    assert loaded.units == []


def test_excel_layout_preserves_business_cells_and_workbook_structure(tmp_path):
    from openpyxl import Workbook
    from openpyxl.worksheet.table import Table

    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Accounts"
    sheet.append(["Account", "Balance"])
    sheet.append(["001", 42])
    sheet["B2"].number_format = "$#,##0.00"
    sheet.merge_cells("A3:B3")
    sheet["A3"] = "Approved"
    sheet.add_table(Table(displayName="Balances", ref="A1:B2"))
    path = tmp_path / "accounts.xlsx"
    workbook.save(path)

    layout = excel_layout(path, document_id="book-1", source_format="xlsx")

    normalized = layout.sheets[0]
    assert layout.service == "openpyxl"
    assert normalized.name == "Accounts"
    assert (normalized.row_count, normalized.col_count) == (3, 2)
    assert normalized.merged_ranges == ["A3:B3"]
    assert normalized.tables == ["Balances"]
    assert normalized.reading_order == ["s0:A1", "s0:B1", "s0:A2", "s0:B2", "s0:A3"]
    balance = next(cell for cell in normalized.cells if cell.ref == "B2")
    assert balance.value == "42"
    assert balance.number_format == "$#,##0.00"
    assert "A2=001 | B2=42" in normalized.content
    assert _col_letter(0) == "A"
    assert _col_letter(26) == "AA"


def test_excel_layout_rejects_corrupt_unsafe_and_unknown_formats(tmp_path):
    corrupt = tmp_path / "corrupt.xlsx"
    corrupt.write_bytes(b"not a zip file")
    with pytest.raises(CorruptFile):
        # CorruptFile is a domain error and intentionally hides parser details.
        inspect_xlsx_safety(corrupt)

    unsafe = tmp_path / "unsafe.xlsx"
    with zipfile.ZipFile(unsafe, "w") as archive:
        archive.writestr("xl/vbaProject.bin", b"macro")
        archive.writestr("xl/externalLinks/externalLink1.xml", b"link")
        archive.writestr("xl/media/payload.exe", b"binary")
    assert inspect_xlsx_safety(unsafe) == ["vba_macros", "external_links", "embedded_objects"]
    with pytest.raises(UnsafeWorkbook):
        excel_layout(unsafe, document_id="unsafe", source_format="xlsx")
    with pytest.raises(ValueError, match="csv"):
        excel_layout(unsafe, document_id="unknown", source_format="csv")
